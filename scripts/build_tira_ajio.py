import sys, csv, datetime as dt, openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils.cell import quote_sheetname as Q

SEP, AUG, MASTER, OUT = sys.argv[1:5]
DUMP, MAP, DD, CV, SOHT = ("Tira Ajio Sales Dump", "TIRA AJIO MAPPER", "Dropdown",
                           "Cat.View Tira Ajio", "SOH Tira Ajio")

def grab(p, s):
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    r = [list(x) for x in wb[s].iter_rows(values_only=True)]; wb.close(); return r
def dd(s): return dt.datetime.strptime(str(s).strip(), "%b %d, %Y").date()

# ---------------- sources ----------------
soh_aug = {str(r[0]).strip(): (str(r[1]).strip(), str(r[3]).strip(), r[4])
           for r in grab(AUG,"SOH")[2:] if r[0]}                       # sap -> (lb, desc, stock)
soh_sep = {str(r[0]).strip(): r[-1] for r in grab(SEP,"SOH")[2:] if r[0]}   # lb -> stock
lb2sap  = {v[0]: k for k, v in soh_aug.items()}
tira_sep, tira_aug = grab(SEP,"Tira Sales")[1:], grab(AUG,"Tira Sales")[1:]
ajio_sep, ajio_aug = grab(SEP,"Ajio Sales")[1:], grab(AUG,"Ajio Sales")[1:]
tira_cat = {str(r[2]).strip(): (r[3], r[5], r[6], r[7]) for r in tira_sep if r[2]}
ajio_cat = {str(r[1]).strip(): (r[2], r[4], r[5]) for r in ajio_sep if r[1]}

master = []
with open(MASTER) as f:
    for sku, sap, name, mrp, land in csv.reader(f, delimiter="\t"):
        master.append((sku.strip(), sap.strip(), name.strip(),
                       float(mrp) if mrp.strip() else None,
                       float(land) if land.strip() else None))

# short name / CATEGORY / TAGS keyed on SAP  (categories = the 13 on your Dropdown tab)
TAX = {
"494836840":("Glam Makeup Box 232g","Make up","Makeup Kit"),
"494836841":("Reetha Shampoo 300ml","Hair","Shampoo"),
"494836842":("Lice Buster Shampoo 40ml","Hair","Shampoo"),
"494836843":("Juicy Locks Hair Oil 100ml","Hair","Hair Oil"),
"494836844":("Tangle Tamer Conditioner 100g","Hair","Conditioner"),
"494836845":("Flake Fighter Shampoo 100ml","Hair","Shampoo"),
"494836846":("Bye Dull Skin Soap 75g x3","Soap","Dull Skin"),
"494836847":("Colourful Soap 3+1 300g","Soap","Combo"),
"494836848":("Tulsi Shower Gel 300ml","Shower gel","Tulsi"),
"494836849":("Oat Lotion 300ml","Lotion","Oat"),
"494836850":("Bye Dull Skin Lotion 100ml","Lotion","Dull Skin"),
"494836851":("Facecream + SPF30 30g","Sun","Face Care"),
"494836852":("Facecream + SPF50 50g","Sun","Face Care"),
"494836853":("Body Lotion + SPF30 100ml","Sun","Lotion"),
"494836854":("Anti Dull Skin Facewash 100ml","Face Care","Dull Skin"),
"494836855":("Kumkumadi Face Cream 50g","Face Care","Kumkumadi"),
"494846201":("Lip Tints Gift Set 4s","Make up","Lip"),
"494846202":("Nothing Lip Balm 4.5g","Make up","Lip"),
"494846203":("Chapstick Cherry 4.5g","Make up","Lip"),
"494846204":("Chapstick Strawberry 4.5g","Make up","Lip"),
"494846205":("Red Beetroot Triple Lip Tint 4.5g","Make up","Lip"),
"494846206":("Kiddy Kohl Kajal 0.3g","Make up","Kajal"),
"494846207":("Daily Facewash 100ml","Face Care","Face Wash"),
"494846208":("Oat Lotion 200ml","Lotion","Oat"),
"494846209":("Oat Bathing Bar 75g","Soap","Oat"),
"494846210":("Kumkumadi Body Lotion 200ml","Lotion","Kumkumadi"),
"494846212":("Alcohol Free Hand Sanitizer 50ml","Sanitizer","Sanitizer"),
"494846213":("Tangle Tamer Spray 50ml","Hair","Conditioner"),
"494846214":("Reetha Shampoo 50ml","Hair","Shampoo"),
"494846215":("Body Wash 50ml","Shower gel","Body Wash"),
"494846216":("Lotion 50ml","Lotion","Others"),
"494846217":("Bubbles & Gum Toothpaste 50g","Toothpaste","Toothpaste"),
}
CATS = ["Hair","Face Care","Dull Skin","Lotion","Make up","Mosquito","Others",
        "Sanitizer","Shower gel","Soap","Sun","Toothpaste","Combo"]
TAGS = sorted({t for _,_,t in TAX.values()} | {"Mosquito","Toothpaste","Others"})
EVTS = ["BAU","1 day sale","P1 sale","Pre buzz","TT"]
order = {s: i for i, (_, s, *_ ) in enumerate(master)}

# ---------------- merge ----------------
rows = []
for src in (tira_sep, tira_aug):
    wide = len(src[0]) > 5
    for r in src:
        if not r[0]: continue
        lb = str(r[2] if wide else r[1]).strip()
        t  = tira_cat.get(lb, ("","","",""))
        rows.append(dict(plat="Tira", date=dd(r[0]), sap=lb2sap.get(lb,""), lb=lb,
            name=(r[3] if wide else r[2]), evt=(str(r[1]).strip() if wide else "BAU"),
            l1=(r[5] if wide else t[1]), l2=(r[6] if wide else t[2]), l3=(r[7] if wide else t[3]),
            qty=float(r[9] if wide else r[4]), gmv=float(r[8] if wide else r[3])))
for src in (ajio_sep, ajio_aug):
    wide = len(src[0]) > 4
    for r in src:
        if not r[0]: continue
        sap = str(r[1]).strip(); a = ajio_cat.get(sap, ("","",""))
        rows.append(dict(plat="Ajio", date=dd(r[0]), sap=sap,
            lb=soh_aug.get(sap,("","",None))[0], name=(r[2] if wide else a[0]), evt="BAU",
            l1=(r[4] if wide else a[1]), l2=(r[5] if wide else a[2]), l3="",
            qty=float(r[7] if wide else r[3]), gmv=float(r[6] if wide else r[2])))
rows.sort(key=lambda x: (x["date"], x["plat"], order.get(x["sap"], 999)))

# ---------------- style ----------------
HDR=PatternFill("solid",fgColor="1F3864"); HDRF=Font(color="FFFFFF",bold=True,size=10)
SUB=PatternFill("solid",fgColor="D9E1F2"); TOT=PatternFill("solid",fgColor="FFF2CC")
INP=PatternFill("solid",fgColor="FCE4D6"); NEW=PatternFill("solid",fgColor="F2F2F2")
TH=Side(style="thin",color="BFBFBF"); BOX=Border(left=TH,right=TH,top=TH,bottom=TH)
DF,MF,IF_,PF = "dd-mmm-yy", '#,##0.00', '#,##0', '0.0%'
def head(ws,h,fr=4):
    for i,v in enumerate(h,1):
        c=ws.cell(row=1,column=i,value=v); c.fill,c.font,c.border=HDR,HDRF,BOX
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.freeze_panes=ws.cell(row=2,column=fr)

wb=Workbook(); wb.remove(wb.active)

# ================= DROPDOWN =================
ws=wb.create_sheet(DD); head(ws,["CATEGORY","TAGS","PLATFORM","EVENT TYPE","GST %"],2)
for i,v in enumerate(CATS,2): ws.cell(row=i,column=1,value=v)
for i,v in enumerate(TAGS,2): ws.cell(row=i,column=2,value=v)
for i,v in enumerate(["Tira","Ajio"],2): ws.cell(row=i,column=3,value=v)
for i,v in enumerate(EVTS,2): ws.cell(row=i,column=4,value=v)
for i,v in enumerate([0.05,0.12,0.18,0.28],2): ws.cell(row=i,column=5,value=v).number_format=PF
for c,w in zip("ABCDE",(18,20,12,14,10)): ws.column_dimensions[c].width=w
N_CAT,N_TAG=1+len(CATS),1+len(TAGS)

# ================= MAPPER =================
ws=wb.create_sheet(MAP)
head(ws,["ARTICLE (SAP)","SKU ID","LB CODE (Tira)","Short Name","Product Name (master)",
         "CATEGORY","TAGS","MRP","Landing Rate","Landing %","COGS W/o GST","GST %","COGS W GST",
         "Tira L1","Tira L2","Tira L3","Ajio Vertical","Ajio Brick","SOH Aug-26","SOH Sep-26"])
for i,(sku,sap,name,mrp,land) in enumerate(master,2):
    lb,desc,st_a = soh_aug.get(sap,("","",None))
    short,cat,tag = TAX.get(sap,("","",""))
    t=tira_cat.get(lb,("","","","")); a=ajio_cat.get(sap,("","",""))
    live = bool(lb)
    V=[sap,sku,lb,short,name,cat,tag,mrp,land,None,None,0.18,None,
       t[1],t[2],t[3],a[1],a[2],st_a,soh_sep.get(lb)]
    for j,v in enumerate(V,1):
        c=ws.cell(row=i,column=j,value=v); c.border=BOX
        if not live: c.fill=NEW
    ws.cell(row=i,column=10,value=f'=IFERROR($I{i}/$H{i},"")').number_format=PF
    ws.cell(row=i,column=13,value=f'=IF($K{i}="","",$K{i}*(1+$L{i}))').number_format=MF
    for cc,fmt in ((8,MF),(9,MF),(11,MF),(12,PF),(19,IF_),(20,IF_)):
        ws.cell(row=i,column=cc).number_format=fmt
    for cc in (6,7,11): ws.cell(row=i,column=cc).fill=INP
LM=1+len(master)
for col,f1,n in (("F",f"={Q(DD)}!$A$2:$A${N_CAT}",0),("G",f"={Q(DD)}!$B$2:$B${N_TAG}",0)):
    dv=DataValidation(type="list",formula1=f1,allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{LM}")
ws.auto_filter.ref=f"A1:T{LM}"
for c,w in zip(["A","B","C","D","E","F","G","H","I","J","K","L","M","N","O","P","Q","R","S","T"],
               (15,9,14,30,42,13,13,9,13,10,13,8,13,13,16,18,13,21,11,11)):
    ws.column_dimensions[c].width=w
ws.cell(row=LM+2,column=1,value="Orange = fill in. COGS W/o GST (col K) is the only true gap; everything downstream is formula-driven.").font=Font(italic=True,color="C00000")
ws.cell(row=LM+3,column=1,value="Grey rows = SKUs with no LB code / no SOH yet (not live on Tira). They map automatically once they start selling.").font=Font(italic=True,color="7F7F7F")

# ================= SOH =================
ws=wb.create_sheet(SOHT); head(ws,["ARTICLE (SAP)","LB CODE","Short Name","SOH Aug-26","SOH Sep-26","Change"],3)
r=2
for sku,sap,name,mrp,land in master:
    lb,desc,st_a=soh_aug.get(sap,("","",None))
    if not lb: continue
    ws.cell(row=r,column=1,value=sap); ws.cell(row=r,column=2,value=lb)
    ws.cell(row=r,column=3,value=TAX.get(sap,("","",""))[0])
    ws.cell(row=r,column=4,value=st_a).number_format=IF_
    ws.cell(row=r,column=5,value=soh_sep.get(lb)).number_format=IF_
    ws.cell(row=r,column=6,value=f'=IFERROR($E{r}-$D{r},"")').number_format=IF_
    for cc in range(1,7): ws.cell(row=r,column=cc).border=BOX
    r+=1
ws.cell(row=r+1,column=3,value="Total").font=Font(bold=True)
for cc in (4,5,6):
    c=ws.cell(row=r+1,column=cc,value=f'=SUM({L(cc)}2:{L(cc)}{r-1})'); c.font=Font(bold=True); c.number_format=IF_
for c,w in zip("ABCDEF",(15,13,30,12,12,11)): ws.column_dimensions[c].width=w

# ================= SALES DUMP =================
ws=wb.create_sheet(DUMP)
head(ws,["PLATFORM","Order Date","ARTICLE (SAP)","LB CODE","SKU ID","Product Name (platform)",
         "Short Name","Event Type","Platform L1","Platform L2","Platform L3",
         "Quantity","GMV","ASP","MRP","Landing Rate","Net Revenue","Platform Margin","Margin %",
         "WEEK","MONTH","CATEGORY","TAGS","COGS","COGS (incl GST)","GM","GM %"])
MR=f"{Q(MAP)}!$A$2:$T${LM}"
def vl(i,col): return f'IFERROR(VLOOKUP($C{i},{MR},{col},0),"")'
for i,r in enumerate(rows,2):
    ws.cell(row=i,column=1,value=r["plat"])
    ws.cell(row=i,column=2,value=r["date"]).number_format=DF
    ws.cell(row=i,column=3,value=r["sap"]); ws.cell(row=i,column=4,value=r["lb"])
    ws.cell(row=i,column=5,value=f'={vl(i,2)}')
    ws.cell(row=i,column=6,value=r["name"]); ws.cell(row=i,column=7,value=f'={vl(i,4)}')
    ws.cell(row=i,column=8,value=r["evt"])
    ws.cell(row=i,column=9,value=r["l1"]); ws.cell(row=i,column=10,value=r["l2"])
    ws.cell(row=i,column=11,value=r["l3"])
    ws.cell(row=i,column=12,value=r["qty"]).number_format=IF_
    ws.cell(row=i,column=13,value=round(r["gmv"],2)).number_format=MF
    for col,f,fmt in (
        (14,f'=IFERROR($M{i}/$L{i},"")',MF), (15,f'={vl(i,8)}',MF), (16,f'={vl(i,9)}',MF),
        (17,f'=IFERROR(IF($P{i}="","",$L{i}*$P{i}),"")',MF),
        (18,f'=IFERROR(IF($Q{i}="","",$M{i}-$Q{i}),"")',MF),
        (19,f'=IFERROR(IF($Q{i}="","",$R{i}/$M{i}),"")',PF),
        (20,f'=$B{i}-WEEKDAY($B{i},3)',DF), (21,f'=TEXT($B{i},"mmm yyyy")',None),
        (22,f'={vl(i,6)}',None), (23,f'={vl(i,7)}',None),
        (24,f'=IFERROR(IF(VLOOKUP($C{i},{MR},11,0)="","",$L{i}*VLOOKUP($C{i},{MR},11,0)),"")',MF),
        (25,f'=IFERROR(IF(VLOOKUP($C{i},{MR},13,0)="","",$L{i}*VLOOKUP($C{i},{MR},13,0)),"")',MF),
        (26,f'=IFERROR(IF(OR($Q{i}="",$Y{i}=""),"",$Q{i}-$Y{i}),"")',MF),
        (27,f'=IFERROR(IF($Z{i}="","",$Z{i}/$Q{i}),"")',PF)):
        c=ws.cell(row=i,column=col,value=f)
        if fmt: c.number_format=fmt
LD=1+len(rows); ws.auto_filter.ref=f"A1:AA{LD}"
for c,w in zip(["A","B","C","D","E","F","G","H","I","J","K","L","M","N","O","P","Q","R","S","T",
                "U","V","W","X","Y","Z","AA"],
               (9,12,15,12,9,32,28,12,14,21,21,9,11,10,9,12,12,14,10,12,11,13,13,11,15,11,8)):
    ws.column_dimensions[c].width=w

# ================= CAT. VIEW =================
ws=wb.create_sheet(CV)
d0,d1=min(r["date"] for r in rows),max(r["date"] for r in rows)
months=[]                      # monthly grain only, derived from the data
for r in sorted(rows,key=lambda x:x["date"]):
    m=r["date"].strftime("%b %Y")
    if m not in months: months.append(m)
weeks,days=[],[]
ws["A1"]="Cat. View — Tira + Ajio"; ws["A1"].font=Font(bold=True,size=13)
ws["B2"]="Platform:"; ws["B2"].font=Font(bold=True)
ws["C2"]="All"; ws["C2"].fill=INP; ws["C2"].font=Font(bold=True)
ws["C2"].alignment=Alignment(horizontal="center")
dvp=DataValidation(type="list",formula1='"All,Tira,Ajio"',allow_blank=False)
ws.add_data_validation(dvp); dvp.add("C2")
ws["A2"]='=IF($C$2="All","*",$C$2)'; ws["A2"].font=Font(color="D9D9D9",size=8)
HR=3
for i,h in enumerate(["Category","Product","Metrics"]+months+weeks+days,1):
    c=ws.cell(row=HR,column=i,value=h); c.fill,c.font,c.border=HDR,HDRF,BOX
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    if i>3+len(months): c.number_format=DF
ws.freeze_panes=ws.cell(row=HR+1,column=4)
CM,CW = 4, 4+len(months); CD=CW+len(weeks); NC=CD+len(days)-1
METRICS=["GMV","QTY","ASP","Net Revenue","COGS (incl GST)","GM","GM %"]
DQ=Q(DUMP)
def block(r0,label,kind,fill=None,bold=False):
    for j,m in enumerate(METRICS):
        r=r0+j
        ws.cell(row=r,column=3,value=m).font=Font(bold=bold,italic=not bold,size=9)
        if j==0:
            ws.cell(row=r,column=2,value=label).font=Font(bold=bold)
        for ci in range(CM,NC+1):
            hc=f"{L(ci)}${HR}"
            dim = "U" if ci<CW else ("T" if ci<CD else "B")
            key = "" if kind is None else (f',{DQ}!$G:$G,$B${r0}' if kind=="P"
                                           else f',{DQ}!$V:$V,$A${r0}')
            S=lambda col: f'SUMIFS({DQ}!${col}:${col},{DQ}!${dim}:${dim},{hc},{DQ}!$A:$A,$A$2{key})'
            if m=="GMV":            f=f'=IFERROR(IF({S("M")}=0,"",{S("M")}),"")'
            elif m=="QTY":          f=f'=IFERROR(IF({S("L")}=0,"",{S("L")}),"")'
            elif m=="ASP":          f=f'=IFERROR(IF({S("L")}=0,"",{S("M")}/{S("L")}),"")'
            elif m=="Net Revenue":  f=f'=IFERROR(IF({S("Q")}=0,"",{S("Q")}),"")'
            elif m=="COGS (incl GST)": f=f'=IFERROR(IF({S("Y")}=0,"",{S("Y")}),"")'
            elif m=="GM":           f=f'=IFERROR(IF({S("Q")}=0,"",{S("Q")}-{S("Y")}),"")'
            else:                   f=f'=IFERROR(IF({S("Q")}=0,"",({S("Q")}-{S("Y")})/{S("Q")}),"")'
            c=ws.cell(row=r,column=ci,value=f)
            c.number_format=IF_ if m=="QTY" else (PF if m=="GM %" else MF)
            if fill: c.fill=fill
            if bold: c.font=Font(bold=True)
        if fill:
            for ci in (1,2,3): ws.cell(row=r,column=ci).fill=fill
    return r0+len(METRICS)
sold={r["sap"] for r in rows}
r=HR+1
ws.cell(row=r,column=1,value="GRAND TOTAL").font=Font(bold=True)
r=block(r,"All products",None,fill=TOT,bold=True)
for cat in [c for c in CATS if any(TAX.get(s,("","",""))[1]==c for s in sold)]:
    mem=[s for s in sorted(sold,key=lambda x:order.get(x,999)) if TAX.get(s,("","",""))[1]==cat]
    ws.cell(row=r,column=1,value=cat).font=Font(bold=True)
    r=block(r,cat,"V",fill=SUB,bold=True)
    for s in mem: r=block(r,TAX[s][0],"P")
LC=r-1
for c,w in zip(["A","B","C"],(14,30,16)): ws.column_dimensions[c].width=w
for ci in range(CM,NC+1): ws.column_dimensions[L(ci)].width=11

wb.save(OUT)
print(f"master SKUs    : {len(master)}   live(with LB/SOH): {sum(1 for _,s,*_ in master if soh_aug.get(s))}")
print(f"merged rows    : {len(rows)}  (Tira {sum(1 for x in rows if x['plat']=='Tira')}, Ajio {sum(1 for x in rows if x['plat']=='Ajio')})")
print(f"unmapped SAP   : {sorted({x['sap'] for x in rows if x['sap'] not in {m[1] for m in master}})}")
print(f"blank SAP rows : {sum(1 for x in rows if not x['sap'])}")
print(f"no landing rate: {sorted({m[1] for m in master if m[4] is None})}")
print(f"dates          : {d0} -> {d1} | months: {months}")
print(f"dump           : rows 2..{LD}, cols A..AA")
print(f"mapper         : rows 2..{LM}")
print(f"cat.view       : rows {HR+1}..{LC}, cols A..{L(NC)}  ({len(sold)} SKUs sold)")
print(f"GMV {sum(x['gmv'] for x in rows):,.2f} | QTY {sum(x['qty'] for x in rows):,.0f}")
